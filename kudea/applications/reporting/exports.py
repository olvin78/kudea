# ============================================================
# EXPORTACIÓN DE INFORMES → CSV / XLSX
# ------------------------------------------------------------
# /reporting/export/ventas/?fmt=xlsx|csv
# /reporting/export/movimientos/?fmt=xlsx|csv
#
# Respeta los mismos filtros de periodo que el dashboard
# (rango, fecha_inicio, fecha_fin) y el mismo criterio de
# ventas (sólo completadas, igual que los KPI).
# ============================================================
import csv
from datetime import datetime, timedelta
from io import BytesIO

from django.http import Http404, HttpResponse
from django.utils.timezone import localdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from applications.cashflow.models import Movimiento
from applications.config.roles import role_required
from applications.home.models import Venta


def resolver_rango(request):
    """(rango, inicio, fin) con la misma lógica que el dashboard."""
    hoy = localdate()
    inicio_mes = hoy.replace(day=1)
    inicio_anio = hoy.replace(month=1, day=1)
    rango = request.GET.get("rango", "hoy")
    f_inicio_str = request.GET.get("fecha_inicio")
    f_fin_str = request.GET.get("fecha_fin")

    if rango == "ayer":
        inicio = fin = hoy - timedelta(days=1)
    elif rango == "semana":
        inicio, fin = hoy - timedelta(days=6), hoy
    elif rango == "mes":
        inicio, fin = inicio_mes, hoy
    elif rango == "anio":
        inicio, fin = inicio_anio, hoy
    elif rango == "personalizado" and f_inicio_str and f_fin_str:
        try:
            inicio = datetime.strptime(f_inicio_str, "%Y-%m-%d").date()
            fin = datetime.strptime(f_fin_str, "%Y-%m-%d").date()
        except ValueError:
            inicio = fin = hoy
    else:
        rango, inicio, fin = "hoy", hoy, hoy
    return rango, inicio, fin


# ------------------------------------------------------------
# Datos
# ------------------------------------------------------------
def _datos_ventas(inicio, fin):
    qs = (
        Venta.objects
        .filter(creado_en__date__gte=inicio, creado_en__date__lte=fin, estado="completada")
        .select_related("cliente", "metodo_pago", "usuario")
        .order_by("creado_en")
    )
    headers = [
        "Codigo", "Fecha", "Hora", "Cliente", "Vendedor", "Metodo",
        "Subtotal", "IVA", "Descuento", "Total", "Devuelto",
    ]
    rows = []
    t_sub = t_iva = t_desc = t_tot = t_dev = 0.0
    for v in qs:
        sub, iva = float(v.subtotal), float(v.iva)
        desc, tot, dev = float(v.descuento), float(v.total), float(v.devuelto_total)
        t_sub += sub; t_iva += iva; t_desc += desc; t_tot += tot; t_dev += dev
        rows.append([
            v.codigo,
            v.creado_en.strftime("%d/%m/%Y"),
            v.creado_en.strftime("%H:%M"),
            v.cliente.nombre if v.cliente_id else "Mostrador",
            v.usuario.username if v.usuario_id else "",
            v.metodo_pago.nombre,
            sub, iva, desc, tot, dev,
        ])
    totales = ["TOTAL", "", "", "", "", "", t_sub, t_iva, t_desc, t_tot, t_dev]
    return headers, rows, totales


def _datos_movimientos(inicio, fin):
    qs = (
        Movimiento.objects
        .filter(fecha__date__gte=inicio, fecha__date__lte=fin)
        .select_related("cuenta", "created_by")
        .order_by("fecha")
    )
    headers = [
        "Fecha", "Hora", "Tipo", "Concepto", "Cuenta", "Metodo",
        "Origen", "Usuario", "Importe",
    ]
    rows = []
    neto = 0.0
    for m in qs:
        importe = float(m.cantidad)
        neto += importe if m.tipo == "ingreso" else -importe
        rows.append([
            m.fecha.strftime("%d/%m/%Y"),
            m.fecha.strftime("%H:%M"),
            m.get_tipo_display(),
            m.concepto,
            m.cuenta.nombre if m.cuenta_id else "",
            m.get_metodo_pago_display(),
            m.get_origen_display(),
            m.created_by.username if m.created_by_id else "",
            importe,
        ])
    totales = ["NETO", "", "", "", "", "", "", "", neto]
    return headers, rows, totales


# ------------------------------------------------------------
# Respuestas
# ------------------------------------------------------------
def _csv_response(nombre, headers, rows, totales):
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.csv"'
    resp.write("\ufeff")  # BOM: que Excel abra bien los acentos
    w = csv.writer(resp, delimiter=";")
    w.writerow(headers)
    w.writerows(rows)
    if totales:
        w.writerow(totales)
    return resp


def _xlsx_response(nombre, headers, rows, totales):
    wb = Workbook()
    ws = wb.active
    ws.title = "Informe"
    ws.append(headers)

    cabecera_fill = PatternFill("solid", fgColor="1D3557")
    cabecera_font = Font(color="FFFFFF", bold=True, size=11)
    borde = Border(bottom=Side(style="thin", color="A8DADC"))
    for cell in ws[1]:
        cell.fill = cabecera_fill
        cell.font = cabecera_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for row in rows:
        ws.append(row)

    if totales:
        ws.append(totales)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True, color="1D3557")
            cell.fill = PatternFill("solid", fgColor="E8F4F6")
            cell.border = borde

    # Formato monetario en las columnas numéricas (desde la 7ª en ventas,
    # la última en movimientos)
    for col_idx, header in enumerate(headers, start=1):
        if header in ("Subtotal", "IVA", "Descuento", "Total", "Devuelto", "Importe"):
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = '#,##0.00'

    # Anchos de columna
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_idx in range(2, min(ws.max_row, 400) + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 45)

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return resp


# ------------------------------------------------------------
# Vista
# ------------------------------------------------------------
@role_required("gerente")
def exportar_informe(request, tipo):
    """Descarga el informe filtrado en XLSX (por defecto) o CSV."""
    if tipo == "ventas":
        headers, rows, totales = _datos_ventas(*resolver_rango(request)[1:])
    elif tipo == "movimientos":
        headers, rows, totales = _datos_movimientos(*resolver_rango(request)[1:])
    else:
        raise Http404

    _, inicio, fin = resolver_rango(request)
    nombre = f"kudea_{tipo}_{inicio:%Y%m%d}_{fin:%Y%m%d}"
    fmt = request.GET.get("fmt", "xlsx").lower()

    if fmt == "csv":
        return _csv_response(nombre, headers, rows, totales)
    return _xlsx_response(nombre, headers, rows, totales)
